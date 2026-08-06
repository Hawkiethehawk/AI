#!/usr/bin/env python3
"""Shared report readers and field mapping for TopOn data exports."""

from __future__ import annotations

import csv
import json
import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable


FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "date": ("日期", "时间", "date", "day", "stat date", "report date"),
    "app": ("应用", "应用名称", "app", "app name", "application", "application name"),
    "app_id": ("应用id", "app id", "application id"),
    "ad_unit": ("广告位", "广告位名称", "placement", "ad unit", "ad placement"),
    "ad_unit_id": ("广告位id", "placement id", "ad unit id", "ad placement id"),
    "country": ("国家", "国家地区", "地区", "country", "geo", "region"),
    "platform": ("平台", "系统", "操作系统", "platform", "os"),
    "app_version": ("应用版本", "版本", "app version", "version"),
    "ad_format": ("广告样式", "广告类型", "广告形式", "ad format", "format"),
    "network": ("广告平台", "平台名称", "network", "ad network", "network name"),
    "ad_source_id": ("广告源id", "ad source id", "adsource id"),
    "ad_source": ("广告源", "广告源名称", "ad source", "ad source name", "adsource"),
    "third_party_ad_unit_id": ("第三方平台广告位id", "三方广告位id", "third party ad unit id", "network placement id"),
    "traffic_group": ("流量分组", "流量组", "traffic group", "segment", "group"),
    "waterfall": ("瀑布流", "瀑布", "waterfall", "mediation group"),
    "source_instance": ("广告源实例", "广告源实例名称", "instance", "source instance"),
    "source_instance_id": ("广告源实例id", "instance id", "source instance id"),
    "layer_order": ("瀑布层级", "排序", "优先级", "layer", "layer order", "priority"),
    "floor_price": ("底价", "排序价格", "floor", "floor price", "sort price"),
    "bidding_type": ("竞价类型", "出价类型", "bidding type", "bid type"),
    "timeout_ms": ("超时时间", "超时毫秒", "timeout", "timeout ms"),
    "latency_ms": ("延迟", "请求耗时", "latency", "latency ms"),
    "config_effective_at": ("配置生效时间", "变更时间", "effective at", "config effective at"),
    "requests": ("广告请求", "请求数", "请求", "requests", "ad requests", "request"),
    "responses": ("广告返回", "返回数", "填充数", "填充请求", "responses", "fills", "filled requests"),
    "impressions": ("展示", "展示数", "展示量", "impressions", "ad impressions"),
    "clicks": ("点击", "点击数", "clicks", "ad clicks"),
    "revenue": ("收益", "收入", "广告收益", "revenue", "income"),
    "estimated_revenue": ("预估收益", "estimated revenue"),
    "ecpm": ("ecpm", "eCPM", "千次展示收益", "千次展示收入"),
    "estimated_ecpm": ("预估 eCPM", "预估eCPM", "estimated ecpm"),
    "fill_rate": ("填充率", "fill rate", "fillrate"),
    "impression_rate": ("展示率", "show rate", "impression rate"),
    "ctr": ("点击率", "ctr", "click through rate"),
    "dau": ("dau", "日活", "日活跃用户", "活跃用户", "daily active users"),
    "ad_users": ("广告用户", "广告用户数", "ad users", "monetized users"),
    "currency": ("币种", "货币", "currency"),
    "timezone": ("时区", "timezone", "time zone"),
    "data_source": ("数据来源", "来源", "data source", "source"),
    "status": ("状态", "status"),
    "statistic_type": ("统计方式", "统计口径", "statistic type", "aggregation type"),
    "deu": ("deu",),
    "new_user_share": ("新用户占比", "new user share"),
    "penetration_rate": ("渗透率", "penetration rate"),
    "estimated_revenue_share": ("预估收益占比", "estimated revenue share"),
    "revenue_api": ("收益 api", "收益api", "revenue api"),
    "revenue_gap_rate": ("收益gap", "revenue gap"),
    "estimated_arpdau": ("预估arpdau", "预估 arpdau", "estimated arpdau"),
    "impressions_per_dau": ("展示/dau", "展示 dau", "impressions per dau"),
    "ad_source_requests": ("广告源请求", "ad source requests"),
    "bidding_requests": ("竞价", "竞价请求", "bidding requests"),
    "bidding_response_ecpm": ("竞价响应ecpm", "bidding response ecpm"),
    "bidding_response_rate": ("竞价响应率", "询价响应率", "竞价询价响应率", "bidding response rate"),
    "bidding_wins": ("竞价胜出数", "竞价胜出次数", "bidding wins", "waterfall wins"),
    "bidding_win_rate": ("竞胜率", "广告竞胜率", "bidding win rate", "waterfall win rate"),
    "ad_source_fill_rate": ("广告源填充率", "ad source fill rate"),
    "sort_price": ("价格", "排序价格", "sort price"),
    "ecpm_api": ("ecpm api", "ecpmapi"),
    "request_or_bid": ("请求/竞价", "请求竞价", "request or bid"),
    "fill_latency_seconds": ("填充耗时（秒）", "填充耗时(秒)", "填充耗时", "fill latency seconds"),
    "click_share": ("点击占比", "click share"),
    "revenue_api_share": ("收益占比", "收益api占比", "revenue api share"),
    "impressions_api_share": ("展示api占比", "impressions api share"),
    "clicks_api_share": ("点击api占比", "clicks api share"),
    "impressions_api": ("展示 api", "展示api", "impressions api"),
    "impression_share": ("展示占比", "impression share"),
    "app_starts": ("应用启动", "app starts"),
    "config_fetches": ("获取配置", "config fetches"),
    "traffic_requests": ("流量请求", "traffic requests"),
    "traffic_fills": ("流量填充", "traffic fills"),
    "traffic_fill_rate": ("流量填充率", "traffic fill rate"),
    "ad_scene_arrivals": ("到达广告场景", "ad scene arrivals"),
    "ad_scene_arrival_rate": ("广告场景到达率", "ad scene arrival rate"),
    "ad_ready_rate": ("广告ready率", "ad ready rate"),
    "is_ready_queries": ("查询isready", "isready queries"),
    "is_ready_success_rate": ("isready成功率", "isready success rate"),
    "display_triggers": ("触发展示", "display triggers"),
    "ad_trigger_rate": ("广告触发率", "ad trigger rate"),
    "display_trigger_successes": ("触发展示成功", "display trigger successes"),
    "display_trigger_success_rate": ("触发展示成功率", "display trigger success rate"),
    "impression_success_rate": ("展示成功率", "impression success rate"),
    "impression_gap_rate": ("展示gap", "impression gap"),
}

ADDITIVE_METRICS = ("requests", "responses", "impressions", "clicks", "revenue")
RATIO_METRICS = ("ecpm", "fill_rate", "show_rate", "ctr")
SUPPLEMENTAL_COUNT_METRICS = (
    "ad_source_requests", "bidding_requests", "revenue_api", "estimated_revenue", "impressions_api",
    "bidding_wins", "request_or_bid",
    "app_starts", "config_fetches", "traffic_requests", "traffic_fills",
    "ad_scene_arrivals", "is_ready_queries", "display_triggers", "display_trigger_successes",
)
SUPPLEMENTAL_RATIO_METRICS = (
    "new_user_share", "penetration_rate", "estimated_revenue_share", "revenue_gap_rate",
    "bidding_response_rate", "ad_source_fill_rate", "impression_share", "traffic_fill_rate",
    "ad_scene_arrival_rate", "ad_ready_rate", "is_ready_success_rate", "ad_trigger_rate",
    "display_trigger_success_rate", "impression_success_rate", "impression_gap_rate", "impression_rate",
    "bidding_win_rate", "click_share", "revenue_api_share", "impressions_api_share", "clicks_api_share",
)
SUPPLEMENTAL_NON_ADDITIVE_METRICS = (
    "deu", "estimated_arpdau", "impressions_per_dau", "bidding_response_ecpm", "estimated_ecpm",
    "sort_price", "ecpm_api", "fill_latency_seconds",
)
DIMENSIONS = (
    "app", "app_id", "ad_unit", "ad_unit_id", "country", "platform",
    "app_version", "ad_format", "network", "traffic_group", "waterfall",
    "source_instance", "source_instance_id", "layer_order", "floor_price",
    "bidding_type", "timeout_ms", "latency_ms", "config_effective_at",
    "ad_source_id", "ad_source", "third_party_ad_unit_id", "status",
    "currency", "timezone", "data_source", "statistic_type",
)
ALL_FIELDS = set(FIELD_ALIASES)


def load_report_type_registry() -> dict[str, Any]:
    path = Path(__file__).resolve().parents[2] / "references" / "topn" / "report-types.json"
    return json.loads(path.read_text(encoding="utf-8"))


def load_metric_catalog() -> dict[str, Any]:
    path = Path(__file__).resolve().parents[2] / "references" / "topn" / "metric-catalog.json"
    return json.loads(path.read_text(encoding="utf-8"))


def metric_definition(field: str) -> dict[str, Any] | None:
    return load_metric_catalog().get("metrics", {}).get(field)


def detect_report_type(headers: Iterable[str], filename: str | None = None) -> dict[str, Any] | None:
    header_set = {normalize_header(item) for item in headers}
    filename_text = normalize_header(filename or "")
    candidates: list[tuple[int, dict[str, Any]]] = []
    for report_type in load_report_type_registry().get("report_types", []):
        required = {normalize_header(item) for item in report_type.get("required_headers", [])}
        if not required.issubset(header_set):
            continue
        signatures = {normalize_header(item) for item in report_type.get("signature_headers", [])}
        signature_score = len(header_set & signatures)
        filename_score = sum(
            10 for pattern in report_type.get("filename_patterns", [])
            if normalize_header(pattern) in filename_text
        )
        candidates.append((len(required) * 100 + signature_score + filename_score, report_type))
    if not candidates:
        return None
    return max(candidates, key=lambda item: item[0])[1]


def is_summary_row(row: dict[str, Any]) -> bool:
    markers = {"合计", "总计", "汇总", "total", "subtotal"}
    for value in row.values():
        if str(value or "").strip().lower() in markers:
            return True
    date_text = str(row.get("date") or row.get("日期") or "").strip()
    return bool(re.match(r"^\d{4}[-/.]\d{2}[-/.]\d{2}\s*[~～至到]\s*\d{4}", date_text))


def normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    return re.sub(r"[\s_\-—–/\\:：()（）\[\]【】]+", "", text)


NORMALIZED_ALIASES = {
    canonical: {normalize_header(alias) for alias in aliases + (canonical,)}
    for canonical, aliases in FIELD_ALIASES.items()
}


def detect_mapping(headers: Iterable[str]) -> tuple[dict[str, str], list[dict[str, str]]]:
    mapping: dict[str, str] = {}
    conflicts: list[dict[str, str]] = []
    for header in headers:
        normalized = normalize_header(header)
        matches = [field for field, aliases in NORMALIZED_ALIASES.items() if normalized in aliases]
        if len(matches) == 1:
            field = matches[0]
            if field in mapping:
                conflicts.append({"canonical": field, "kept": mapping[field], "ignored": header})
            else:
                mapping[field] = header
    return mapping, conflicts


def validate_mapping(mapping: dict[str, str], headers: Iterable[str]) -> list[dict[str, str]]:
    errors: list[dict[str, str]] = []
    header_set = set(headers)
    used_sources: dict[str, str] = {}
    for field, source in mapping.items():
        if field not in ALL_FIELDS:
            errors.append({"code": "unknown_canonical_field", "field": field, "source": source})
        if source not in header_set:
            errors.append({"code": "source_column_not_found", "field": field, "source": source})
        if source in used_sources and used_sources[source] != field:
            errors.append({
                "code": "source_column_reused", "field": field,
                "source": source, "also_mapped_to": used_sources[source],
            })
        used_sources[source] = field
    return errors


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, Any]], dict[str, Any]]:
    raw = path.read_bytes()
    encoding = None
    text = None
    for candidate in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            text = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if text is None or encoding is None:
        raise ValueError("Unable to decode CSV as UTF-8 or GB18030")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(text.splitlines(), dialect=dialect)
    headers = [str(item or "").strip() for item in (reader.fieldnames or [])]
    rows = [{str(key or "").strip(): value for key, value in row.items()} for row in reader]
    return headers, rows, {"format": "csv", "encoding": encoding, "delimiter": dialect.delimiter}


def _read_xlsx(path: Path, sheet: str | None) -> tuple[list[str], list[dict[str, Any]], dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Reading XLSX requires openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_name = sheet or workbook.sheetnames[0]
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}. Available: {', '.join(workbook.sheetnames)}")
    worksheet = workbook[sheet_name]
    iterator = worksheet.iter_rows(values_only=True)
    try:
        first_row = next(iterator)
    except StopIteration:
        return [], [], {"format": "xlsx", "sheet": sheet_name, "sheets": workbook.sheetnames}
    headers = [str(value or "").strip() for value in first_row]
    rows: list[dict[str, Any]] = []
    for values in iterator:
        if not any(value not in (None, "") for value in values):
            continue
        rows.append({headers[index]: values[index] if index < len(values) else None for index in range(len(headers))})
    return headers, rows, {"format": "xlsx", "sheet": sheet_name, "sheets": workbook.sheetnames}


def read_report(path: Path, sheet: str | None = None) -> tuple[list[str], list[dict[str, Any]], dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv", ".txt"}:
        return _read_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(path, sheet)
    raise ValueError(f"Unsupported report format: {suffix}. Use CSV, TSV, XLSX or XLSM.")


def parse_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = str(value).strip()
    if not text or text.lower() in {"nan", "n/a", "null", "none", "--", "-"}:
        return None
    percentage = text.endswith("%")
    text = text.rstrip("%").replace(",", "")
    text = re.sub(r"^[¥￥$€£]", "", text).strip()
    try:
        number = float(text)
    except ValueError:
        return None
    return number / 100 if percentage else number


def parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        return None


def parse_date_range_from_filename(filename: str) -> dict[str, str] | None:
    """Extract a YYYYMMDD-YYYYMMDD or YYYY-MM-DD_YYYY-MM-DD range."""
    matches = re.findall(r"(?<!\d)(20\d{2})[-_.]?(\d{2})[-_.]?(\d{2})(?!\d)", filename)
    if len(matches) < 2:
        return None
    parsed: list[date] = []
    for year, month, day in matches[:2]:
        try:
            parsed.append(date(int(year), int(month), int(day)))
        except ValueError:
            return None
    if parsed[0] > parsed[1]:
        return None
    return {"start": parsed[0].isoformat(), "end": parsed[1].isoformat()}


def load_mapping(value: str | None) -> dict[str, str]:
    if not value:
        return {}
    candidate = Path(value)
    try:
        is_file = candidate.exists()
    except OSError:
        is_file = False
    payload = candidate.read_text(encoding="utf-8") if is_file else value
    parsed = json.loads(payload)
    if not isinstance(parsed, dict):
        raise ValueError("Mapping must be a JSON object of canonical_field -> source_column")
    return {str(key): str(column) for key, column in parsed.items()}


def canonical_rows(rows: list[dict[str, Any]], mapping: dict[str, str]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for row in rows:
        item: dict[str, Any] = {}
        for field, source in mapping.items():
            value = row.get(source)
            if field == "date":
                parsed = parse_date(value)
                item[field] = parsed.isoformat() if parsed else None
            elif field == "sort_price":
                parsed = parse_number(value)
                item[field] = parsed if parsed is not None else (None if value is None else str(value).strip())
            elif field in (
                ADDITIVE_METRICS + RATIO_METRICS + SUPPLEMENTAL_COUNT_METRICS
                + SUPPLEMENTAL_RATIO_METRICS + SUPPLEMENTAL_NON_ADDITIVE_METRICS
                + ("dau", "ad_users")
            ):
                item[field] = parse_number(value)
            else:
                item[field] = None if value is None else str(value).strip()
        output.append(item)
    return output


def write_json(payload: Any, output: Path | None) -> None:
    text = json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n"
    if output:
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(text, encoding="utf-8", newline="\n")
    else:
        print(text, end="")
