# -*- coding: utf-8 -*-
import datetime
import re
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def current_monday():
    d = datetime.datetime.utcnow().date()
    return (d - datetime.timedelta(days=d.weekday())).isoformat()


HEAD_FILL = PatternFill('solid', fgColor='1F3864')
HEAD_FONT = Font(bold=True, color='FFFFFF', size=10)
TITLE_FONT = Font(bold=True, size=14, color='1F3864')
NOTE_FONT = Font(size=9, color='595959')
LINK_FONT = Font(color='0563C1', underline='single', size=10)
LINK_FONT_BOLD = Font(color='0563C1', underline='single', size=10, bold=True)
UP_FILL = PatternFill('solid', fgColor='E2EFDA')
DOWN_FILL = PatternFill('solid', fgColor='FCE4E4')
NEW_FILL = PatternFill('solid', fgColor='FFF2CC')
FOCUS_FILL = PatternFill('solid', fgColor='F8CBAD')
thin = Side(style='thin', color='D9D9D9')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
GROUP_TOP = Border(left=thin, right=thin, bottom=thin, top=Side(style='medium', color='1F3864'))
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)

COLS = [
    ('序号', 5, 'c'),
    ('游戏名', 28, 'l'),
    ('品类', 12, 'c'),
    ('Tag路径', 22, 'l'),
    ('本周排名', 7, 'c'),
    ('上周排名', 7, 'c'),
    ('变化量', 7, 'c'),
    ('6周排名轨迹', 30, 'l'),
    ('Top50稳定性', 20, 'l'),
    ('近30天下载国Top5', 30, 'l'),
    ('近30天收入国Top5', 26, 'l'),
    ('市场属性', 20, 'l'),
    ('上线日期', 12, 'c'),
    ('评分', 6, 'c'),
    ('评论数', 11, 'c'),
    ('发行商', 20, 'l'),
    ('总部', 6, 'l'),
    ('重点关注', 8, 'c'),
    ('备注', 26, 'l'),
]


def build_column_index(cols=COLS):
    return {name: i + 1 for i, (name, _, _) in enumerate(cols)}


def chain_taxonomy(tax):
    byid = {t['id']: t for t in tax}
    parents = set(p for t in tax for p in (t.get('parent_ids') or []))
    leaves = [t for t in tax if t['id'] not in parents] or tax
    best = []
    for lf in leaves:
        chain, cur, seen = [], lf, set()
        while cur and cur['id'] not in seen:
            seen.add(cur['id'])
            chain.append(cur['name'])
            pid = (cur.get('parent_ids') or [None])[0]
            cur = byid.get(pid)
        chain = list(reversed(chain))
        if len(chain) > len(best):
            best = chain
    return best


def tag_path(tags):
    tags = tags or []
    domain = next((t['name'] for t in tags if t.get('type') == 'domain'), '')
    games = [t for t in tags if t.get('type') == 'games']
    seq = [domain] if domain else []
    if games:
        meta = next((t['name'] for t in tags if t.get('type') == 'meta'), '')
        if meta:
            seq.append(meta)
        seq += chain_taxonomy(games)
    else:
        seq += chain_taxonomy([t for t in tags if t.get('type') == 'apps'])
    out = []
    for x in seq:
        if x and x not in out:
            out.append(x)
    return ' / '.join(out[:3])


def norm_date(s):
    if not s:
        return ''
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})', s)
    return f'{m.group(1)}-{m.group(2)}-{m.group(3)}' if m else s


def hist_str(history):
    return ' → '.join(str(x) if x is not None else '·' for x in reversed(history))


def stability(r):
    history = r['history']
    count50 = sum(1 for h in history if h is not None and h <= 50)
    onboard = sum(1 for h in history if h is not None)
    parts = []
    if onboard > 0:
        parts.append(f'在榜{onboard}周')
    if count50 > 0:
        parts.append(f'Top50 {count50}周')
    return '/'.join(parts)


def change_str(r):
    if r['lastWeek'] is None:
        return 'NEW'
    c = r['change']
    return f'+{c}' if c > 0 else (str(c) if c < 0 else '0')


def top5(s):
    if not s:
        return ''
    return ' / '.join([x for x in s.split(' / ') if x][:5])


def store_url(store_ids):
    sids = store_ids or []
    gp = next((s for s in sids if s.startswith('1_')), None)
    if gp:
        return 'https://play.google.com/store/apps/details?id=' + gp[2:]
    ios = next((s for s in sids if s.startswith('2_')), None) or next((s for s in sids if s.startswith('3_')), None)
    if ios:
        return 'https://apps.apple.com/app/id' + ios[2:]
    return ''


def build_focus_cells(r, cat):
    co = r.get('country') or {}
    reasons = '/'.join(r.get('_focusReasons', []))
    return {
        '序号': None,
        '游戏名': r.get('name', ''),
        '品类': cat,
        'Tag路径': tag_path(r.get('tags')),
        '本周排名': r['rank'],
        '上周排名': r['lastWeek'] if r['lastWeek'] is not None else 'NEW',
        '变化量': change_str(r),
        '6周排名轨迹': hist_str(r['history']),
        'Top50稳定性': stability(r),
        '近30天下载国Top5': top5(co.get('dlList', '')) or '—',
        '近30天收入国Top5': top5(co.get('revList', '')) or '—（无内购收入/IAA变现）',
        '市场属性': co.get('market', ''),
        '上线日期': norm_date(r.get('release', '')),
        '评分': round(r['rating'], 2) if r.get('rating') else '—',
        '评论数': f"{r['reviews']:,}" if r.get('reviews') else '—',
        '发行商': r.get('publisher', ''),
        '总部': r.get('hq', ''),
        '重点关注': '是' if r.get('_focus', False) else '',
        '备注': reasons if reasons else '',
    }


def apply_focus_row(ws, row, seq, r, cat, col_map, cols=COLS, *, first_in_group=False, category_fill=None):
    cells = build_focus_cells(r, cat)
    cells['序号'] = seq
    border = GROUP_TOP if first_in_group else BORDER
    for name, _, align in cols:
        j = col_map[name]
        cell = ws.cell(row, j, cells[name])
        cell.border = border
        cell.alignment = CENTER if align == 'c' else LEFT

    name_cell = ws.cell(row, col_map['游戏名'])
    url = store_url(r.get('storeIds'))
    if url:
      name_cell.hyperlink = url
      name_cell.font = LINK_FONT_BOLD if r.get('_focus', False) else LINK_FONT

    change_cell = ws.cell(row, col_map['变化量'])
    if r['lastWeek'] is None:
        change_cell.fill = NEW_FILL
    elif r['change'] is not None and r['change'] > 0:
        change_cell.fill = UP_FILL
    elif r['change'] is not None and r['change'] < 0:
        change_cell.fill = DOWN_FILL

    if category_fill is not None:
        ws.cell(row, col_map['品类']).fill = category_fill

    if r.get('_focus', False):
        name_cell.fill = FOCUS_FILL
